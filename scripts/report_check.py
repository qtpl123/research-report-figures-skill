#!/usr/bin/env python
"""Word 报告版面与图片质检（在交付前跑）。

用法:
    python report_check.py <report.docx> [--manifest assets/manifest.json] [--json out.json]

检查项:
1. 图片是否全部为“嵌入式”（浮动图片容易错位/遮挡文字）
2. 每张图片的显示尺寸是否超出正文区域（页边距内）
3. 图片高度是否超过单页可容纳高度（避免跨页/遮挡）
4. 每张图片后面是否紧跟图注（以“图 N”开头）
5. 图注是否包含来源/许可说明
6. 表格是否超出版心宽度
7. 文档是否包含目录域
"""

from __future__ import annotations

import argparse
import json
import sys
import zipfile
from pathlib import Path

try:  # 避免在 GBK 控制台下因非 GBK 字符报错
    sys.stdout.reconfigure(errors="replace")
except Exception:
    pass

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu

EMU_PER_CM = 360000


def content_box_cm(doc: Document) -> tuple[float, float]:
    s = doc.sections[0]
    w = (s.page_width - s.left_margin - s.right_margin) / EMU_PER_CM
    h = (s.page_height - s.top_margin - s.bottom_margin) / EMU_PER_CM
    return w, h


def has_toc_field(docx_path: Path) -> bool:
    with zipfile.ZipFile(docx_path) as z:
        xml = z.read("word/document.xml").decode("utf-8", "ignore")
    return "TOC \\o" in xml or "TOC \\\\o" in xml or 'TOC \\o "1-3"' in xml or " TOC " in xml


def has_update_fields(docx_path: Path) -> bool:
    with zipfile.ZipFile(docx_path) as z:
        try:
            xml = z.read("word/settings.xml").decode("utf-8", "ignore")
        except KeyError:
            return False
    return "updateFields" in xml


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--manifest", default=None)
    ap.add_argument("--json", default=None)
    args = ap.parse_args()

    docx_path = Path(args.docx).resolve()
    doc = Document(str(docx_path))
    box_w, box_h = content_box_cm(doc)
    max_img_h = box_h * 0.78

    # 逐个段落扫描 drawing
    drawings = []
    for idx, p in enumerate(doc.paragraphs):
        for drawing in p._p.findall(".//" + qn("w:drawing")):
            inline = drawing.find(qn("wp:inline"))
            anchor = drawing.find(qn("wp:anchor"))
            ext = (inline if inline is not None else anchor)
            extent = ext.find(qn("wp:extent")) if ext is not None else None
            cx = int(extent.get("cx")) if extent is not None else 0
            cy = int(extent.get("cy")) if extent is not None else 0
            drawings.append(
                {
                    "paragraph": idx,
                    "floating": anchor is not None,
                    "width_cm": cx / EMU_PER_CM,
                    "height_cm": cy / EMU_PER_CM,
                    "following": (doc.paragraphs[idx + 1].text.strip()
                                  if idx + 1 < len(doc.paragraphs) else ""),
                }
            )

    issues: list[str] = []
    warnings: list[str] = []

    for i, d in enumerate(drawings, 1):
        if d["floating"]:
            issues.append(f"图 {i}: 使用了浮动定位（可能遮挡文字），应改为嵌入式")
        if d["width_cm"] > box_w + 0.05:
            issues.append(
                f"图 {i}: 宽度 {d['width_cm']:.2f}cm 超出版心 {box_w:.2f}cm"
            )
        if d["height_cm"] > max_img_h:
            warnings.append(
                f"图 {i}: 高度 {d['height_cm']:.2f}cm 偏大（建议 ≤ {max_img_h:.2f}cm），可能跨页"
            )
        nxt = d["following"]
        if not nxt.startswith("图"):
            issues.append(f"图 {i}: 后面不是图注（实际内容: {nxt[:30]!r}）")
        else:
            src = doc.paragraphs[d["paragraph"] + 2].text.strip()
            if not any(k in src for k in ("来源", "链接", "许可", "本图", "非真实")):
                warnings.append(f"图 {i}: 未找到来源/许可说明行")

    # 表格宽度
    for ti, table in enumerate(doc.tables, 1):
        try:
            widths = [c.width for c in table.rows[0].cells if c.width is not None]
            total = sum(widths) / EMU_PER_CM if widths else 0
            if total > box_w + 0.1:
                issues.append(f"表 {ti}: 合计宽度 {total:.2f}cm 超出版心 {box_w:.2f}cm")
        except Exception:
            pass

    if not has_toc_field(docx_path):
        warnings.append("未检测到目录域（TOC）")
    if not has_update_fields(docx_path):
        warnings.append("未设置“打开时更新域”，目录/页码可能显示为空白")

    report = {
        "docx": str(docx_path),
        "content_box_cm": {"width": round(box_w, 2), "height": round(box_h, 2)},
        "images": len(drawings),
        "tables": len(doc.tables),
        "paragraphs": len(doc.paragraphs),
        "issues": issues,
        "warnings": warnings,
    }

    print(f"报告质检：{docx_path.name}")
    print(f"  正文版心 {box_w:.2f}cm × {box_h:.2f}cm；图片 {len(drawings)} 张，表格 {len(doc.tables)} 个，段落 {len(doc.paragraphs)} 段")
    if issues:
        print(f"  [需修复] 问题 {len(issues)} 项：")
        for x in issues:
            print(f"    - {x}")
    else:
        print("  [通过] 无阻塞性问题")
    if warnings:
        print(f"  [提醒] {len(warnings)} 项：")
        for x in warnings:
            print(f"    - {x}")

    if args.json:
        Path(args.json).write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
        )

    return 1 if issues else 0


if __name__ == "__main__":
    raise SystemExit(main())
